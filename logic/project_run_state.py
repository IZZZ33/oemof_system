"""Persistent project-input fingerprints and simulation run status."""

from __future__ import annotations

import hashlib
import json
import math
import os
import struct
from datetime import date, datetime, time, timezone
from pathlib import Path

from logic.project_paths import get_project_dir, get_project_results_dir, get_results_root


_IGNORED_INPUT_FILES = {
    "input_done.flag",
    "simulation_status.json",
    # Durable download cache. The weather values that affect a simulation are
    # already stored in each scenario workbook and are fingerprinted there.
    "dwd_weather.json",
    # Files created by operating systems and file browsers are not model inputs.
    ".DS_Store",
    "Thumbs.db",
    "desktop.ini",
}

_TEXT_INPUT_SUFFIXES = {".csv", ".tsv", ".txt"}
_JSON_INPUT_SUFFIXES = {".json", ".geojson"}
_SEMANTIC_DIGEST_CACHE: dict[str, tuple[int, int, bytes]] = {}


def _is_transient_input_file(path: Path) -> bool:
    """Return whether *path* is an editor/office artefact, not a model input."""
    name = path.name
    return (
        name in _IGNORED_INPUT_FILES
        or name.startswith("~$")  # Microsoft Office owner/lock file
        or name.endswith((".tmp", ".temp", ".swp"))
    )


def _input_files(root: Path):
    return sorted(
        (p for p in root.rglob("*") if p.is_file() and not _is_transient_input_file(p)),
        key=lambda p: p.relative_to(root).as_posix(),
    )


def _update_tagged(digest, tag: bytes, value: bytes) -> None:
    digest.update(tag)
    digest.update(len(value).to_bytes(8, "big"))
    digest.update(value)


def _cell_value_bytes(value) -> bytes:
    """Encode an Excel cell value without workbook/ZIP metadata."""
    if isinstance(value, bool):
        return b"bool:true" if value else b"bool:false"
    if isinstance(value, int):
        return b"int:" + str(value).encode("ascii")
    if isinstance(value, float):
        if math.isnan(value):
            return b"float:nan"
        if math.isinf(value):
            return b"float:+inf" if value > 0 else b"float:-inf"
        return b"float:" + struct.pack(">d", value)
    if isinstance(value, (datetime, date, time)):
        return b"datetime:" + value.isoformat().encode("utf-8")
    if isinstance(value, bytes):
        return b"bytes:" + value
    if isinstance(value, str):
        value = value.replace("\r\n", "\n").replace("\r", "\n")
        return b"str:" + value.encode("utf-8")
    return (f"{type(value).__name__}:{value}").encode("utf-8")


def _xlsx_content_digest(path: Path) -> bytes:
    """Hash workbook names, formulas, and cell values, but not Excel metadata."""
    from openpyxl import load_workbook

    digest = hashlib.sha256()
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            _update_tagged(digest, b"sheet", worksheet.title.encode("utf-8"))
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    _update_tagged(digest, b"cell", cell.coordinate.encode("ascii"))
                    _update_tagged(digest, b"value", _cell_value_bytes(cell.value))
    finally:
        workbook.close()
    return digest.digest()


def _file_content_digest(path: Path) -> bytes:
    """Return a portable digest of the input content that affects the model."""
    stat = path.stat()
    cache_key = str(path.resolve())
    cached = _SEMANTIC_DIGEST_CACHE.get(cache_key)
    if cached and cached[:2] == (stat.st_mtime_ns, stat.st_size):
        return cached[2]

    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        result = _xlsx_content_digest(path)
    elif suffix in _JSON_INPUT_SUFFIXES:
        try:
            with path.open("r", encoding="utf-8-sig") as stream:
                value = json.load(stream)
            canonical = json.dumps(
                value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ).encode("utf-8")
            result = hashlib.sha256(canonical).digest()
        except (UnicodeDecodeError, json.JSONDecodeError):
            result = _raw_file_digest(path)
    elif suffix in _TEXT_INPUT_SUFFIXES:
        try:
            text = path.read_text(encoding="utf-8-sig")
            canonical = text.replace("\r\n", "\n").replace("\r", "\n")
            result = hashlib.sha256(canonical.encode("utf-8")).digest()
        except UnicodeDecodeError:
            result = _raw_file_digest(path)
    else:
        result = _raw_file_digest(path)

    _SEMANTIC_DIGEST_CACHE[cache_key] = (stat.st_mtime_ns, stat.st_size, result)
    return result


def _raw_file_digest(path: Path) -> bytes:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.digest()


def project_input_fingerprint(project_name: str) -> str:
    """Hash model-relevant input content in a machine-independent form.

    Excel ZIP metadata, workbook formatting, line-ending differences, and
    temporary Office/OS files are intentionally excluded. Actual cell values,
    formulas, filenames, and other durable input content remain significant.
    """
    root = Path(get_project_dir(project_name))
    digest = hashlib.sha256()
    if not root.exists():
        return digest.hexdigest()

    for path in _input_files(root):
        relative = path.relative_to(root).as_posix().encode("utf-8")
        _update_tagged(digest, b"path", relative)
        _update_tagged(digest, b"content", _file_content_digest(path))
    return digest.hexdigest()


def _legacy_project_input_fingerprint(project_name: str) -> str:
    """Reproduce the former raw-byte hash for already simulated projects."""
    root = Path(get_project_dir(project_name))
    digest = hashlib.sha256()
    if not root.exists():
        return digest.hexdigest()

    for path in _input_files(root):
        relative = path.relative_to(root).as_posix().encode("utf-8")
        digest.update(len(relative).to_bytes(4, "big"))
        digest.update(relative)
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
    return digest.hexdigest()


def run_status_path(project_name: str) -> Path:
    return Path(get_project_results_dir(project_name)) / "simulation_status.json"


def read_run_status(project_name: str) -> dict:
    path = run_status_path(project_name)
    try:
        with path.open("r", encoding="utf-8") as stream:
            data = json.load(stream)
        if not isinstance(data, dict):
            return {}
        if data.get("state") == "running" and not _process_is_alive(data.get("runner_pid")):
            write_run_status(
                project_name,
                "interrupted",
                input_fingerprint=data.get("input_fingerprint"),
                scenario=data.get("scenario"),
                scenario_index=data.get("scenario_index"),
                scenario_count=data.get("scenario_count"),
                interrupted_from="running",
                error="The previous simulation process stopped before completion.",
            )
            return read_run_status(project_name)
        return data
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return {}


def write_run_status(project_name: str, state: str, **details) -> None:
    path = run_status_path(project_name)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "project_name": project_name,
        "state": state,
        "updated_at": datetime.now(timezone.utc).isoformat(),
        **details,
    }
    temporary = path.with_suffix(".tmp")
    with temporary.open("w", encoding="utf-8") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2)
    temporary.replace(path)


def _process_is_alive(pid) -> bool:
    try:
        process_id = int(pid)
        if process_id <= 0:
            return False
        if os.name == "nt":
            # os.kill(pid, 0) can raise SystemError with some Windows Python
            # builds. Query the process handle without sending any signal.
            import ctypes

            process_query_limited_information = 0x1000
            still_active = 259
            kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
            open_process = kernel32.OpenProcess
            open_process.argtypes = [ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32]
            open_process.restype = ctypes.c_void_p
            get_exit_code = kernel32.GetExitCodeProcess
            get_exit_code.argtypes = [ctypes.c_void_p, ctypes.POINTER(ctypes.c_uint32)]
            get_exit_code.restype = ctypes.c_int
            close_handle = kernel32.CloseHandle
            close_handle.argtypes = [ctypes.c_void_p]
            close_handle.restype = ctypes.c_int

            handle = open_process(
                process_query_limited_information, False, process_id
            )
            if not handle:
                # Access denied means that the process exists but cannot be
                # inspected by this user. Other errors mean it no longer exists.
                return ctypes.get_last_error() == 5
            try:
                exit_code = ctypes.c_uint32()
                return bool(get_exit_code(handle, ctypes.byref(exit_code))) and (
                    exit_code.value == still_active
                )
            finally:
                close_handle(handle)

        os.kill(process_id, 0)
        return True
    except (TypeError, ValueError, OSError, PermissionError, SystemError):
        return False


def recover_interrupted_run_statuses(
    pending_projects=(), results_root: str | Path | None = None
) -> list[str]:
    """Convert orphaned queued/running states into recoverable interruptions.

    This is called once when the simulation runner starts. A queued project is
    retained only when its global submission flag still exists. A running state
    is retained only when the recorded runner process is still alive.
    """
    pending = {str(project) for project in (pending_projects or ()) if project}
    root = Path(results_root) if results_root is not None else Path(get_results_root())
    recovered = []
    if not root.exists():
        return recovered

    for status_path in root.glob("*/simulation_status.json"):
        try:
            with status_path.open("r", encoding="utf-8") as stream:
                status = json.load(stream)
        except (OSError, json.JSONDecodeError):
            continue
        if not isinstance(status, dict):
            continue

        project_name = str(status.get("project_name") or status_path.parent.name)
        state = status.get("state")
        if state == "queued" and project_name in pending:
            continue
        if state == "running" and _process_is_alive(status.get("runner_pid")):
            continue
        if state not in {"queued", "running"}:
            continue

        write_root = run_status_path(project_name)
        if results_root is not None:
            write_root = root / project_name / "simulation_status.json"
            write_root.parent.mkdir(parents=True, exist_ok=True)
            payload = {
                **status,
                "project_name": project_name,
                "state": "interrupted",
                "updated_at": datetime.now(timezone.utc).isoformat(),
                "interrupted_from": state,
                "error": "The previous simulation process stopped before completion.",
            }
            temporary = write_root.with_suffix(".tmp")
            with temporary.open("w", encoding="utf-8") as stream:
                json.dump(payload, stream, ensure_ascii=False, indent=2)
            temporary.replace(write_root)
        else:
            write_run_status(
                project_name,
                "interrupted",
                input_fingerprint=status.get("input_fingerprint"),
                scenario=status.get("scenario"),
                scenario_index=status.get("scenario_index"),
                scenario_count=status.get("scenario_count"),
                interrupted_from=state,
                error="The previous simulation process stopped before completion.",
            )
        recovered.append(project_name)
    return recovered


def results_match_current_inputs(project_name: str, status: dict | None = None) -> bool:
    status = status or read_run_status(project_name)
    completed_fingerprint = status.get("input_fingerprint") if status.get("state") == "completed" else None
    if not completed_fingerprint:
        return False
    if completed_fingerprint == project_input_fingerprint(project_name):
        return True
    # Status files written before portable fingerprints were introduced contain
    # a raw-byte hash. Keep those results valid as long as the old inputs really
    # are unchanged; transient Excel and OS files are still ignored.
    return completed_fingerprint == _legacy_project_input_fingerprint(project_name)
